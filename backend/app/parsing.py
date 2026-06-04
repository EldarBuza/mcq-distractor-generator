"""Parse user-supplied questions from CSV, JSON, or pasted text.

Each parser returns a ParseResult: the questions it could build, plus a list of
human-readable per-row errors for anything it had to skip. Nothing here calls
the API — it only normalizes input into validated QuestionInput objects.
"""
import csv
import io
import json

from pydantic import BaseModel, ValidationError

from app.models import Difficulty, QuestionInput

# Accepted header aliases (lowercased) -> canonical field.
_QUESTION_KEYS = {"question", "questions", "q", "prompt", "stem"}
_ANSWER_KEYS = {"correct_answer", "answer", "correct", "a", "key"}
_NUM_KEYS = {"num_distractors", "distractors", "n", "num"}
_DIFFICULTY_KEYS = {"difficulty", "level"}


class ParseResult(BaseModel):
    questions: list[QuestionInput]
    errors: list[str]


def _coerce(
    question: str | None,
    answer: str | None,
    num: object = None,
    difficulty: object = None,
) -> QuestionInput:
    """Build a validated QuestionInput, applying defaults for optional fields."""
    data: dict = {
        "question": (question or "").strip(),
        "correct_answer": (answer or "").strip(),
    }
    if num not in (None, ""):
        data["num_distractors"] = int(num)
    if difficulty not in (None, ""):
        data["difficulty"] = Difficulty(str(difficulty).strip().lower())
    return QuestionInput(**data)


def _pick(row: dict, keys: set[str]) -> object:
    for k, v in row.items():
        if k is not None and k.strip().lower() in keys:
            return v
    return None


def parse_csv(text: str) -> ParseResult:
    """Parse CSV with a header row. Recognizes flexible column-name aliases."""
    questions: list[QuestionInput] = []
    errors: list[str] = []

    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        return ParseResult(questions=[], errors=["CSV is empty or has no header row."])

    lowered = {f.strip().lower() for f in reader.fieldnames if f}
    if not (_QUESTION_KEYS & lowered) or not (_ANSWER_KEYS & lowered):
        errors.append(
            "CSV must have a question column (e.g. 'question') and an answer "
            "column (e.g. 'correct_answer')."
        )
        return ParseResult(questions=[], errors=errors)

    for i, row in enumerate(reader, start=2):  # row 1 is the header
        try:
            q = _coerce(
                _pick(row, _QUESTION_KEYS),
                _pick(row, _ANSWER_KEYS),
                _pick(row, _NUM_KEYS),
                _pick(row, _DIFFICULTY_KEYS),
            )
            questions.append(q)
        except (ValidationError, ValueError) as exc:
            errors.append(f"Row {i}: {_short(exc)}")

    return ParseResult(questions=questions, errors=errors)


def parse_json(text: str) -> ParseResult:
    """Parse JSON: either a list of objects or {"questions": [...]}."""
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        return ParseResult(questions=[], errors=[f"Invalid JSON: {exc}"])

    if isinstance(data, dict) and "questions" in data:
        data = data["questions"]
    if not isinstance(data, list):
        return ParseResult(
            questions=[],
            errors=["JSON must be a list of questions or {\"questions\": [...]}."],
        )

    questions: list[QuestionInput] = []
    errors: list[str] = []
    for i, item in enumerate(data, start=1):
        if not isinstance(item, dict):
            errors.append(f"Item {i}: expected an object, got {type(item).__name__}.")
            continue
        try:
            q = _coerce(
                item.get("question") or _first(item, _QUESTION_KEYS),
                item.get("correct_answer") or _first(item, _ANSWER_KEYS),
                item.get("num_distractors") or _first(item, _NUM_KEYS),
                item.get("difficulty") or _first(item, _DIFFICULTY_KEYS),
            )
            questions.append(q)
        except (ValidationError, ValueError) as exc:
            errors.append(f"Item {i}: {_short(exc)}")

    return ParseResult(questions=questions, errors=errors)


def parse_pasted_text(text: str, delimiter: str | None = None) -> ParseResult:
    """Parse pasted lines of 'question <sep> answer'.

    The separator is auto-detected per the first data line (tab, '|', then ';'),
    or pass `delimiter` to force one. Lines without a separator are reported.
    """
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if not lines:
        return ParseResult(questions=[], errors=["No non-empty lines to parse."])

    sep = delimiter or _detect_sep(lines[0])
    if sep is None:
        return ParseResult(
            questions=[],
            errors=[
                "Could not detect a separator. Put one question per line as "
                "'question | answer' (or use a tab or ';')."
            ],
        )

    questions: list[QuestionInput] = []
    errors: list[str] = []
    for i, line in enumerate(lines, start=1):
        parts = line.split(sep)
        if len(parts) < 2:
            errors.append(f"Line {i}: no '{sep}' separator found.")
            continue
        q_text, answer = parts[0], parts[1]
        num = parts[2] if len(parts) > 2 else None
        difficulty = parts[3] if len(parts) > 3 else None
        try:
            questions.append(_coerce(q_text, answer, num, difficulty))
        except (ValidationError, ValueError) as exc:
            errors.append(f"Line {i}: {_short(exc)}")

    return ParseResult(questions=questions, errors=errors)


def parse_text(text: str, fmt: str = "auto") -> ParseResult:
    """Parse a raw text blob (e.g. pasted input) in the given format.

    fmt: "auto" | "csv" | "json" | "lines". "auto" sniffs the content: JSON if it
    starts with [ or {, CSV if the first line looks like a header with the
    required columns, otherwise line-delimited 'question | answer'.
    """
    if fmt == "json":
        return parse_json(text)
    if fmt == "csv":
        return parse_csv(text)
    if fmt == "lines":
        return parse_pasted_text(text)

    # auto
    stripped = text.lstrip()
    if stripped.startswith("[") or stripped.startswith("{"):
        return parse_json(text)
    first_line = next((ln for ln in text.splitlines() if ln.strip()), "")
    header = {c.strip().lower() for c in first_line.split(",")}
    if (_QUESTION_KEYS & header) and (_ANSWER_KEYS & header):
        return parse_csv(text)
    return parse_pasted_text(text)


def parse_upload(filename: str, content: bytes) -> ParseResult:
    """Dispatch by file extension to the right parser.

    .docx and .pdf are binary: their text is extracted first, then run through
    the same auto-detecting text parser (so a document containing 'q | a' lines,
    a CSV-style block, JSON, or a two-column Q/A table is understood).
    """
    name = (filename or "").lower()

    if name.endswith(".docx"):
        text, err = _extract_docx(content)
        return ParseResult(questions=[], errors=[err]) if err else parse_text(text)
    if name.endswith(".pdf"):
        text, err = _extract_pdf(content)
        return ParseResult(questions=[], errors=[err]) if err else parse_text(text)

    try:
        text = content.decode("utf-8-sig")  # tolerate a BOM
    except UnicodeDecodeError:
        return ParseResult(questions=[], errors=["File is not valid UTF-8 text."])

    if name.endswith(".json"):
        return parse_json(text)
    if name.endswith(".csv") or name.endswith(".tsv") or name.endswith(".txt"):
        return parse_csv(text)
    # Unknown extension: best-effort — try JSON, fall back to CSV.
    stripped = text.lstrip()
    if stripped.startswith("[") or stripped.startswith("{"):
        return parse_json(text)
    return parse_csv(text)


def _extract_docx(content: bytes) -> tuple[str, str | None]:
    """Extract text from a .docx. Paragraphs become lines; each table row becomes
    a '|'-joined line so a two-column Question/Answer table parses directly.
    Returns (text, error)."""
    try:
        from docx import Document
    except ImportError:  # pragma: no cover - dependency is in requirements
        return "", "Reading .docx files requires the python-docx package."
    try:
        doc = Document(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001 - surface a clean message to the user
        return "", f"Could not read .docx file: {_short(exc)}"

    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines), None


def _extract_pdf(content: bytes) -> tuple[str, str | None]:
    """Extract text from a .pdf, page by page. Returns (text, error)."""
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover - dependency is in requirements
        return "", "Reading .pdf files requires the pypdf package."
    try:
        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:  # noqa: BLE001 - surface a clean message to the user
        return "", f"Could not read .pdf file: {_short(exc)}"
    if not text.strip():
        return "", (
            "No text found in the PDF (it may be scanned images rather than "
            "selectable text)."
        )
    return text, None


# ---- helpers ----------------------------------------------------------------


def _detect_sep(line: str) -> str | None:
    for sep in ("\t", "|", ";"):
        if sep in line:
            return sep
    return None


def _first(item: dict, keys: set[str]) -> object:
    for k, v in item.items():
        if k.lower() in keys:
            return v
    return None


def _short(exc: Exception) -> str:
    """Compress a pydantic error to a one-line message."""
    if isinstance(exc, ValidationError):
        parts = []
        for e in exc.errors():
            loc = ".".join(str(x) for x in e.get("loc", ())) or "input"
            parts.append(f"{loc}: {e.get('msg', 'invalid')}")
        return "; ".join(parts)
    return str(exc)
