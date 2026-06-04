from app.models import Difficulty
from app.parsing import (
    parse_csv,
    parse_json,
    parse_pasted_text,
    parse_text,
    parse_upload,
)


# ---- CSV --------------------------------------------------------------------


def test_csv_basic():
    text = "question,correct_answer\nCapital of Japan?,Tokyo\n2+2?,4\n"
    res = parse_csv(text)
    assert len(res.questions) == 2
    assert res.errors == []
    assert res.questions[0].question == "Capital of Japan?"
    assert res.questions[0].correct_answer == "Tokyo"
    assert res.questions[0].num_distractors == 3  # default


def test_csv_with_optional_columns_and_aliases():
    text = "Q,Answer,n,level\nCapital of Japan?,Tokyo,4,hard\n"
    res = parse_csv(text)
    assert len(res.questions) == 1
    q = res.questions[0]
    assert q.num_distractors == 4
    assert q.difficulty == Difficulty.hard


def test_csv_missing_required_column():
    text = "foo,bar\n1,2\n"
    res = parse_csv(text)
    assert res.questions == []
    assert any("question column" in e for e in res.errors)


def test_csv_reports_bad_row_but_keeps_good_ones():
    text = "question,correct_answer\nGood?,Yes\n,MissingQ\n"
    res = parse_csv(text)
    assert len(res.questions) == 1
    assert len(res.errors) == 1
    assert res.errors[0].startswith("Row 3")


# ---- JSON -------------------------------------------------------------------


def test_json_list():
    text = '[{"question": "Q?", "correct_answer": "A", "difficulty": "easy"}]'
    res = parse_json(text)
    assert len(res.questions) == 1
    assert res.questions[0].difficulty == Difficulty.easy


def test_json_wrapped_object():
    text = '{"questions": [{"question": "Q?", "correct_answer": "A"}]}'
    res = parse_json(text)
    assert len(res.questions) == 1


def test_json_invalid():
    res = parse_json("{not json")
    assert res.questions == []
    assert any("Invalid JSON" in e for e in res.errors)


# ---- pasted text ------------------------------------------------------------


def test_paste_pipe_delimited():
    res = parse_pasted_text("Capital of Japan? | Tokyo\nWho wrote Hamlet? | Shakespeare")
    assert len(res.questions) == 2
    assert res.questions[1].correct_answer == "Shakespeare"


def test_paste_tab_delimited_with_options():
    res = parse_pasted_text("Q?\tA\t5\thard")
    assert len(res.questions) == 1
    assert res.questions[0].num_distractors == 5
    assert res.questions[0].difficulty == Difficulty.hard


def test_paste_no_separator():
    res = parse_pasted_text("just a line with no separator")
    assert res.questions == []
    assert res.errors


# ---- dispatch ---------------------------------------------------------------


def test_parse_upload_by_extension():
    csv_res = parse_upload("q.csv", b"question,correct_answer\nQ?,A\n")
    assert len(csv_res.questions) == 1
    json_res = parse_upload("q.json", b'[{"question":"Q?","correct_answer":"A"}]')
    assert len(json_res.questions) == 1


def test_parse_upload_handles_bom():
    res = parse_upload("q.csv", "﻿question,correct_answer\nQ?,A\n".encode("utf-8"))
    assert len(res.questions) == 1


# ---- .docx / .pdf uploads ---------------------------------------------------


def _build_docx(paragraphs: list[str], table_rows: list[tuple[str, str]]) -> bytes:
    """Build a real .docx in memory for round-trip testing."""
    import io

    from docx import Document

    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        table = doc.add_table(rows=0, cols=2)
        for q, a in table_rows:
            cells = table.add_row().cells
            cells[0].text = q
            cells[1].text = a
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_parse_upload_docx_pipe_paragraphs():
    content = _build_docx(["Capital of Japan? | Tokyo", "2 + 2? | 4"], [])
    res = parse_upload("quiz.docx", content)
    assert len(res.questions) == 2
    assert res.questions[0].question == "Capital of Japan?"
    assert res.questions[0].correct_answer == "Tokyo"


def test_parse_upload_docx_two_column_table():
    content = _build_docx([], [("Capital of France?", "Paris"), ("3 x 3?", "9")])
    res = parse_upload("quiz.docx", content)
    assert len(res.questions) == 2
    answers = {q.correct_answer for q in res.questions}
    assert {"Paris", "9"} == answers


def test_parse_upload_docx_corrupt_returns_clean_error():
    res = parse_upload("broken.docx", b"not a real docx")
    assert res.questions == []
    assert res.errors and "docx" in res.errors[0].lower()


def test_parse_upload_pdf_corrupt_returns_clean_error():
    res = parse_upload("broken.pdf", b"%PDF-not-really")
    assert res.questions == []
    assert res.errors and "pdf" in res.errors[0].lower()


# ---- parse_text auto-detection ----------------------------------------------


def test_parse_text_auto_json():
    res = parse_text('[{"question":"Q?","correct_answer":"A"}]')
    assert len(res.questions) == 1


def test_parse_text_auto_csv():
    res = parse_text("question,correct_answer\nQ?,A\n")
    assert len(res.questions) == 1


def test_parse_text_auto_lines():
    res = parse_text("Capital of Japan? | Tokyo")
    assert len(res.questions) == 1
    assert res.questions[0].correct_answer == "Tokyo"


def test_parse_text_explicit_format_overrides_sniff():
    # Looks like it could sniff as lines, but force csv.
    res = parse_text("question,correct_answer\nQ?,A\n", fmt="csv")
    assert len(res.questions) == 1
